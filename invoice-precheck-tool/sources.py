# -*- coding: utf-8 -*-
"""Data sources: billing-app queue, reference workbook, NetSuite (pluggable)."""
import json
import re
import urllib.request
from openpyxl import load_workbook

BILLING_APP = "http://192.168.0.1:3000"
WORKBOOK = r"C:\Users\<user>\Desktop\invoice审查系统\invoice-pre-check-reference.xlsx"
OVERSEAS_SUFFIX = re.compile(r"-\s*[A-Za-z]")


def fetch_billing_queue(timeout=20):
    """Team3 + Unbilled orders from the billing app. Raises on failure (fail closed)."""
    with urllib.request.urlopen(BILLING_APP + "/db", timeout=timeout) as r:
        db = json.load(r)
    queue = [o for o in db.get("orders", [])
             if o.get("workflow_stage") == "Team3"
             and o.get("billing_status", "Unbilled") == "Unbilled"]
    seen, out = set(), []
    for o in queue:
        if o["so_number"] not in seen:
            seen.add(o["so_number"])
            out.append(o)
    return out


def load_reference_workbook(path=WORKBOOK):
    """Load and validate the reference workbook. Returns dict or raises ValueError."""
    wb = load_workbook(path, data_only=True)
    need = ["Customers", "Seasonal Discounts", "Special Discounts", "Rules"]
    missing = [s for s in need if s not in wb.sheetnames]
    if missing:
        raise ValueError(f"工作簿缺少工作表：{missing}")
    meta = {}
    for s in need:
        ws = wb[s]
        status = str(ws["B1"].value or "").strip()
        if status != "ACTIVE":
            raise ValueError(f"工作表『{s}』的 Data Status 是『{status}』，不是 ACTIVE — 请先启用")
        meta[s] = {"schema": str(ws["B2"].value or ""), "updated_at": str(ws["B3"].value or ""),
                   "updated_by": str(ws["B4"].value or "")}

    rules = {}
    for row in wb["Rules"].iter_rows(min_row=8, max_col=2):
        k = row[0].value
        if k:
            rules[str(k).strip()] = row[1].value
    for req in ("MAX_ITEM_RATE_USD", "DISCOUNT_AMOUNT_TOLERANCE_USD", "ROUNDING_MODE",
                "ROUNDING_SCALE", "ALLOWED_SO_STATUSES", "ALLOWED_RESULT_STATUSES",
                "SUPPORTED_SCHEMA_VERSION"):
        if req not in rules:
            raise ValueError(f"Rules 工作表缺少必填参数 {req}")
    rules["ALLOWED_SO_STATUSES"] = [s.strip() for s in str(rules["ALLOWED_SO_STATUSES"]).split("|")]
    supported = str(rules["SUPPORTED_SCHEMA_VERSION"]).strip()
    for s, m in meta.items():
        if m["schema"].strip() != supported:
            raise ValueError(f"工作表『{s}』的 schema 版本 {m['schema']} 与工具支持的 {supported} 不一致")

    # 一个客户内部 ID 可以有多行 —— 不同生效期（方案 §4.2）。所以这里是
    # id -> [行, …]，由 engine.select_customer_row 按 SO 日期挑出适用的那行。
    # 旧代码在这里对重复 ID 直接 raise，等于"给某个客户加一个新折扣期"就让
    # 整个工作簿加载失败、当天全队列跑不了。生效期有问题现在只影响该客户。
    customers = {}
    for row in wb["Customers"].iter_rows(min_row=8, max_col=9):
        iid = row[0].value
        if not iid:
            continue
        iid = str(iid).strip()
        name = str(row[2].value or "").strip()
        c = {
            "internal_id": iid, "code": str(row[1].value or "").strip(), "name": name,
            "handling": str(row[3].value or "STANDARD").strip(),
            "can_rule": str(row[4].value or "DEFAULT").strip(),
            "effective_from": row[5].value,
            "effective_to": row[6].value,
            "active": str(row[7].value or "").strip().upper() == "TRUE",
            "notes": str(row[8].value or "").strip(),
        }
        suffix = bool(OVERSEAS_SUFFIX.search(name))
        c["overseas"] = (c["can_rule"] == "FORCE-CAN") or (suffix and c["can_rule"] != "NOT-CAN")
        customers.setdefault(iid, []).append(c)

    def rows(sheet, ncols):
        out = []
        for row in wb[sheet].iter_rows(min_row=8, max_col=ncols):
            if any(c.value not in (None, "") for c in row):
                out.append([c.value for c in row])
        return out

    return {"meta": meta, "rules": rules, "customers": customers,
            "seasonal_rows": rows("Seasonal Discounts", 12),
            "special_rows": rows("Special Discounts", 16)}


class NetSuiteSource:
    """Interface for fetching normalized SO + customer data.

    Implementations must return the dict shapes engine.run_checks expects.
    - PlaywrightSource (current): drives a persistent logged-in browser profile.
    - RestApiSource (later): NetSuite REST/SuiteQL with an integration token.
      Swapping implementations must not require engine changes.
    """

    def fetch_so(self, so_number):
        raise NotImplementedError

    def fetch_customer(self, internal_id):
        raise NotImplementedError


# Real implementation lives in netsuite_scraper.py (xml=T over an
# authenticated persistent-profile session). Re-exported here so app.py and
# the future RestApiSource share one import site.
from netsuite_scraper import PlaywrightSource, NeedLoginError  # noqa: E402,F401
